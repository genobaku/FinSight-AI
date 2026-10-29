"""
DOCX Parser for Microsoft Word documents.
Extracts headings, paragraphs, embedded tables, and document metadata.
"""

import io
from typing import Tuple, List, Dict, Any
from docx import Document
from backend.document_processing.models import ExtractedTable, ExtractedField, SourceEvidence

class DocxParser:
    @staticmethod
    def parse(file_bytes: bytes, filename: str) -> Tuple[str, List[ExtractedTable], List[ExtractedField], Dict[str, Any]]:
        tables: List[ExtractedTable] = []
        fields: List[ExtractedField] = []
        metadata: Dict[str, Any] = {"format": "docx"}
        raw_text_parts: List[str] = []

        try:
            doc = Document(io.BytesIO(file_bytes))

            # 1. Core Properties
            props = doc.core_properties
            if props.title:
                fields.append(ExtractedField(
                    id="field-doc-title",
                    label="Document Title",
                    value=props.title,
                    data_type="text",
                    semantic_category="title",
                    source_evidence=SourceEvidence(text=props.title, page=1),
                    confidence=1.0,
                    verification_status="verified"
                ))
            if props.author:
                fields.append(ExtractedField(
                    id="field-doc-author",
                    label="Author",
                    value=props.author,
                    data_type="text",
                    semantic_category="identity",
                    source_evidence=SourceEvidence(text=props.author, page=1),
                    confidence=1.0,
                    verification_status="verified"
                ))

            # 2. Paragraphs & Headings
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                raw_text_parts.append(text)

                if p.style.name.startswith("Heading"):
                    raw_text_parts.append(f"## {text}")

            # 3. Embedded Tables
            for t_idx, t in enumerate(doc.tables):
                if not t.rows:
                    continue

                headers = [cell.text.strip() for cell in t.rows[0].cells]
                rows = []
                for r in t.rows[1:]:
                    row_vals = [cell.text.strip() for cell in r.cells]
                    if any(row_vals):
                        rows.append(row_vals)

                table_id = f"table-docx-{t_idx+1}"
                tables.append(ExtractedTable(
                    id=table_id,
                    title=f"Table {t_idx+1}",
                    headers=headers,
                    rows=rows,
                    source_evidence={"table_index": t_idx+1, "total_rows": len(rows)}
                ))

                raw_text_parts.append(f"--- Table {t_idx+1} ---")
                raw_text_parts.append(" | ".join(headers))
                for r in rows[:15]:
                    raw_text_parts.append(" | ".join(r))

            metadata["paragraphs_count"] = len(doc.paragraphs)
            metadata["table_count"] = len(doc.tables)

        except Exception as e:
            raise ValueError(f"Failed to parse Word (.docx) document: {str(e)}")

        full_raw_text = "\n".join(raw_text_parts)
        return full_raw_text, tables, fields, metadata
